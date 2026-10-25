import sys
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import linprog
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ------------------------------------------------------------------ 1. DOC DU LIEU
def clean_sku(s):
    return str(s).strip().upper()


def load_data(folder):
    folder = Path(folder)
    product = pd.read_excel(folder / "Master_Data.xlsx", sheet_name="Product_BOM")
    factory = pd.read_excel(folder / "Master_Data.xlsx", sheet_name="Factory_Capacity")
    stock = pd.read_excel(folder / "Inventory_System.xlsx", sheet_name="Stock_Level")
    demand = pd.read_csv(folder / "Demand_Forecast.csv")
    trans = pd.read_csv(folder / "Transport_Costs.csv")

    for df, col in [(product, "SKU"), (stock, "SKU"), (demand, "SKU")]:
        df[col] = df[col].map(clean_sku)
    demand["Region"] = demand["Region"].astype(str).str.strip().str.title()
    trans["Destination"] = trans["Destination"].astype(str).str.strip().str.title()
    trans["Source"] = trans["Source"].astype(str).str.strip().str.upper()
    factory["Factory_ID"] = factory["Factory_ID"].astype(str).str.strip().str.upper()

    # Cau truc du lieu: dictionary
    P = {
        "I": list(factory["Factory_ID"]),
        "J": sorted(demand["Region"].unique()),
        "K": list(product["SKU"]),
        "Cap": dict(zip(factory["Factory_ID"], factory["Công suất tối đa (giờ)"])),
        "Cprod": dict(zip(factory["Factory_ID"], factory["Chi phí vận hành ($/giờ)"])),
        "T": dict(zip(product["SKU"], product["Công lao động (giờ/sp)"])),
        "M": dict(zip(product["SKU"], product["Định mức Vải (m/sp)"])),
        "Q": dict(zip(stock["SKU"], stock["Tồn kho hiện tại"])),
        "S": dict(zip(stock["SKU"], stock["Mức an toàn"])),
        "D": {(r.Region, r.SKU): r.Forecasted_Demand for r in demand.itertuples()},
        "Ctrans": {(r.Source, r.Destination): r.Unit_Transport_Cost for r in trans.itertuples()},
        "Name": dict(zip(product["SKU"], product["Tên Sản Phẩm"])),
        "Loc": dict(zip(factory["Factory_ID"], factory["Địa điểm"])),
    }
    return P

def solve(P):
    I, J, K = P["I"], P["J"], P["K"]
    idx = [(i, j, k) for i in I for j in J for k in K]      
    n = len(idx)                                            

    c = [P["T"][k] * P["Cprod"][i] + P["Ctrans"][(i, j)] for i, j, k in idx] + [0.0]

    A_ub, b_ub, A_eq, b_eq = [], [], [], []

    # R1: ton kho an toan (dang <=  nen doi dau)
    for k in K:
        A_ub.append([-1.0 if kk == k else 0.0 for _, _, kk in idx] + [0.0])
        b_ub.append(P["Q"][k] - sum(P["D"][(j, k)] for j in J) - P["S"][k])

    # R3: nhu cau tung vung
    for j in J:
        for k in K:
            A_ub.append([-1.0 if (jj == j and kk == k) else 0.0 for _, jj, kk in idx] + [0.0])
            b_ub.append(-P["D"][(j, k)])

    # R2: can bang tai
    for i in I:
        A_eq.append([P["T"][k] if ii == i else 0.0 for ii, _, k in idx] + [-P["Cap"][i]])
        b_eq.append(0.0)

    bounds = [(0, None)] * n + [(0, 1)]
    res = linprog(c, A_ub=A_ub, b_ub=b_ub, A_eq=A_eq, b_eq=b_eq,
                  bounds=bounds, method="highs")
    if res.status != 0:
        raise RuntimeError(f"Khong tim duoc nghiem toi uu: {res.message}")

    X = {idx[t]: float(res.x[t]) for t in range(n)}
    return X, float(res.x[-1]), float(res.fun)


# ------------------------------------------------------------------ 3. TONG HOP KET QUA
def build_tables(P, X, rho, Z, tol=1e-6):
    I, J, K = P["I"], P["J"], P["K"]

    plan = pd.DataFrame(
        [[i, k, P["Name"][k], sum(X[(i, j, k)] for j in J)] for i in I for k in K],
        columns=["Nhà máy", "SKU", "Tên sản phẩm", "Số lượng sản xuất"])

    routes = pd.DataFrame(
        [[i, j, k, P["Name"][k], X[(i, j, k)], P["Ctrans"][(i, j)],
          X[(i, j, k)] * P["Ctrans"][(i, j)]]
         for i in I for j in J for k in K if X[(i, j, k)] > tol],
        columns=["Từ nhà máy", "Đến DC", "SKU", "Tên sản phẩm", "Số lượng",
                 "Cước/sp ($)", "Chi phí vận chuyển ($)"])

    util_rows = []
    for i in I:
        hours = sum(P["T"][k] * X[(i, j, k)] for j in J for k in K)
        prod_cost = sum(P["T"][k] * P["Cprod"][i] * X[(i, j, k)] for j in J for k in K)
        trans_cost = sum(P["Ctrans"][(i, j)] * X[(i, j, k)] for j in J for k in K)
        fabric = sum(P["M"][k] * X[(i, j, k)] for j in J for k in K)
        util_rows.append([i, P["Loc"][i], P["Cap"][i], hours, hours / P["Cap"][i],
                          prod_cost, trans_cost, fabric])
    util = pd.DataFrame(util_rows, columns=[
        "Nhà máy", "Địa điểm", "Công suất (giờ)", "Giờ sử dụng", "Tỷ lệ sử dụng",
        "Chi phí sản xuất ($)", "Chi phí vận chuyển ($)", "Vải sử dụng (m)"])

    check_rows = []
    for j in J:
        for k in K:
            got = sum(X[(i, j, k)] for i in I)
            check_rows.append([j, k, P["D"][(j, k)], got, got - P["D"][(j, k)]])
    demand_chk = pd.DataFrame(check_rows, columns=["DC", "SKU", "Nhu cầu", "Giao đến", "Chênh lệch"])

    stock_rows = []
    for k in K:
        end = P["Q"][k] + sum(X[(i, j, k)] for i in I for j in J) - sum(P["D"][(j, k)] for j in J)
        stock_rows.append([k, P["Q"][k], sum(X[(i, j, k)] for i in I for j in J),
                           sum(P["D"][(j, k)] for j in J), end, P["S"][k], end - P["S"][k]])
    stock = pd.DataFrame(stock_rows, columns=[
        "SKU", "Tồn hiện tại", "Sản xuất", "Nhu cầu", "Tồn cuối kỳ", "Mức an toàn", "Dư so với an toàn"])

    total_prod = float(util["Chi phí sản xuất ($)"].sum())
    total_trans = float(util["Chi phí vận chuyển ($)"].sum())
    summary = pd.DataFrame([
        ["Tổng chi phí tối ưu Z ($)", Z],
        ["  - Chi phí sản xuất ($)", total_prod],
        ["  - Chi phí vận chuyển ($)", total_trans],
        ["Tỷ lệ sử dụng công suất chung ρ", rho],
        ["Tổng giờ công sử dụng", float(util["Giờ sử dụng"].sum())],
        ["Tổng vải sử dụng (m)", float(util["Vải sử dụng (m)"].sum())],
        ["Tổng sản lượng (sp)", float(plan["Số lượng sản xuất"].sum())],
    ], columns=["Chỉ tiêu", "Giá trị"])
    return summary, plan, routes, util, demand_chk, stock


# ------------------------------------------------------------------ 4. XUAT EXCEL
def export_excel(path, tables, names):
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for df, name in zip(tables, names):
            df.to_excel(w, sheet_name=name, index=False)
        thin = Side(style="thin", color="999999")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        head_fill = PatternFill("solid", fgColor="1F4E78")
        for ws in w.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in ws.iter_rows():                    # ke bang bang vong For
                for cell in row:
                    cell.border = border
                    if isinstance(cell.value, float):
                        cell.number_format = "0.0%" if ("sử dụng" in str(ws.cell(1, cell.column).value)
                                                        and "Tỷ lệ" in str(ws.cell(1, cell.column).value)) \
                            else "#,##0.0"
            for col in range(1, ws.max_column + 1):
                width = max(len(str(c.value)) if c.value is not None else 0
                            for c in ws[get_column_letter(col)])
                ws.column_dimensions[get_column_letter(col)].width = min(max(width + 3, 12), 42)
        # dinh dang rieng cho ρ o sheet Tong quan
        ws = w.book["Tổng quan"]
        for r in range(2, ws.max_row + 1):
            if "ρ" in str(ws.cell(r, 1).value):
                ws.cell(r, 2).number_format = "0.0%"


# ------------------------------------------------------------------ 5. MAIN
def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else "data"
    out = sys.argv[2] if len(sys.argv) > 2 else "Ket_qua_toi_uu.xlsx"

    P = load_data(folder)
    X, rho, Z = solve(P)
    tables = build_tables(P, X, rho, Z)
    names = ["Tổng quan", "Kế hoạch sản xuất", "Lộ trình vận chuyển",
             "Sử dụng nhà máy", "Kiểm tra nhu cầu", "Kiểm tra tồn kho"]
    export_excel(out, tables, names)

    print(f"Z = {Z:,.2f} $   rho = {rho:.4f}")
    print(tables[3].to_string(index=False))
    print(f"Da xuat ket qua: {out}")


if __name__ == "__main__":
    main()
